import os
import sqlite3
import json
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BASE_DIR, "data", "nusantarabus.db")
SENTIMENT_PATH = os.path.join(BASE_DIR, "data", "sentiment_summary.json")
RAW_SCRAPED_PATH = os.path.join(BASE_DIR, "data", "raw_scraped_schedules.json")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "data", "nusantarabus_complete_dataset.xlsx")

print(f"Reading data from {DB_PATH}, {SENTIMENT_PATH}, and {RAW_SCRAPED_PATH}...")

conn = sqlite3.connect(DB_PATH)

# 1. Sheet 1: Ulasan Riil Granular Per Penumpang (Siapa yang komen & apa komentarnya)
query_reviews = """
SELECT 
    reviewer_name AS "Nama Penumpang (Reviewer)",
    source_platform AS "Sumber Platform",
    review_date AS "Tanggal Ulasan",
    operator_name AS "Operator Bus (PO)",
    route_traveled AS "Rute Perjalanan",
    service_class AS "Kelas Layanan",
    rating_stars AS "Rating (1-5)",
    sentiment_label AS "Sentimen",
    aspect_highlight AS "Aspek Utama",
    review_title AS "Judul Ulasan",
    review_comment AS "Komentar Lengkap Penumpang"
FROM fact_passenger_reviews
ORDER BY operator_name, review_date DESC;
"""
df_granular_reviews = pd.read_sql_query(query_reviews, conn)

# 2. Sheet 2: Ringkasan Sentimen 10 PO Bus
with open(SENTIMENT_PATH, "r", encoding="utf-8") as f:
    sentiment_data = json.load(f)

sentiment_rows = []
for po, s in sentiment_data.items():
    aspects = s.get("aspect_scores", {})
    quotes = " | ".join(s.get("verified_quotes", []))
    sentiment_rows.append({
        "Nama PO Bus": po,
        "Skor Kepuasan (1-5)": s.get("rating_score"),
        "Total Ulasan Terverifikasi": s.get("total_reviews"),
        "Sentimen Positif (%)": s.get("sentiment_positive_pct"),
        "Sentimen Netral (%)": s.get("sentiment_neutral_pct"),
        "Sentimen Negatif (%)": s.get("sentiment_negative_pct"),
        "Label Sentimen": s.get("sentiment_label"),
        "Kenyamanan Kursi (1-5)": aspects.get("comfort_sleeper"),
        "Kebersihan Kabin & Toilet (1-5)": aspects.get("cabin_cleanliness"),
        "Pelayanan Kru Bus (1-5)": aspects.get("crew_hospitality"),
        "Kualitas Makanan Servis (1-5)": aspects.get("food_service"),
        "Ketepatan Waktu Lapangan (1-5)": aspects.get("punctuality"),
        "Keunggulan Utama": s.get("top_strengths"),
        "Poin Perhatian Penumpang": s.get("top_considerations")
    })
df_sentiment_summary = pd.DataFrame(sentiment_rows)

# 3. Sheet 3: Jadwal Scraped RedBus
with open(RAW_SCRAPED_PATH, "r", encoding="utf-8") as f:
    raw_scraped = json.load(f)

scraped_list = raw_scraped.get("schedules", [])
df_raw_scraped = pd.DataFrame(scraped_list)
if not df_raw_scraped.empty:
    df_raw_scraped.rename(columns={
        "operator": "Operator Bus",
        "route": "Kode Rute",
        "service_class": "Kelas Layanan",
        "departure": "Jam Berangkat",
        "arrival": "Jam Tiba",
        "origin": "Titik Naik / Terminal Asal",
        "destination": "Titik Turun / Terminal Tujuan",
        "price": "Harga Tiket Riil (Rp)",
        "rating": "Rating Publik (1-5)"
    }, inplace=True)

# 4. Sheet 4: Peringkat Kinerja PO
query_operators = """
SELECT 
    o.operator_id AS "ID Operator",
    o.po_name AS "Nama Operator Bus",
    o.service_class AS "Kelas Layanan Utama",
    o.headquarters AS "Hub / Kantor Pusat",
    o.fleet_count AS "Estimasi Jumlah Armada",
    COUNT(t.trip_id) AS "Total Trip Diaudit",
    ROUND(AVG(t.arrival_delay_min), 1) AS "Rata-Rata Delay (Menit)",
    ROUND(SUM(CASE WHEN t.otp_status = 'ON_TIME' THEN 1 ELSE 0 END) * 100.0 / COUNT(t.trip_id), 1) AS "Rasio Tepat Waktu OTP (%)",
    ROUND(SUM(CASE WHEN t.otp_status = 'SEVERE_DELAY' THEN 1 ELSE 0 END) * 100.0 / COUNT(t.trip_id), 1) AS "Risiko Delay Parah >60m (%)",
    ROUND(AVG(t.avg_speed_kmh), 1) AS "Kecepatan Rata-Rata (km/jam)",
    ROUND(AVG(t.fuel_intensity_l_100km), 2) AS "Konsumsi Solar (L/100km)"
FROM dim_operators o
LEFT JOIN fact_trips t ON o.operator_id = t.operator_id
GROUP BY o.operator_id
ORDER BY "Rasio Tepat Waktu OTP (%)" DESC;
"""
df_operators = pd.read_sql_query(query_operators, conn)

# 5. Sheet 5: Sampel Perjalanan Aktual (500 Trip)
query_trips = """
SELECT 
    t.trip_id AS "ID Perjalanan",
    o.po_name AS "Operator Bus",
    r.origin_city || ' (' || r.origin_terminal || ')' AS "Asal Keberangkatan",
    r.destination_city || ' (' || r.destination_terminal || ')' AS "Tujuan Kedatangan",
    t.scheduled_departure_ts AS "Jadwal Berangkat",
    t.actual_departure_ts AS "Aktual Berangkat",
    t.scheduled_arrival_ts AS "Jadwal Tiba",
    t.actual_arrival_ts AS "Aktual Tiba",
    ROUND(t.arrival_delay_min, 1) AS "Keterlambatan (Menit)",
    t.otp_status AS "Status OTP",
    t.ticket_price_idr AS "Tarif Tiket (IDR)",
    t.distance_km AS "Jarak (km)",
    ROUND(t.fuel_consumed_liters, 1) AS "Solar Terpakai (L)",
    ROUND(t.avg_speed_kmh, 1) AS "Kecepatan Tol (km/j)",
    t.passenger_count AS "Penumpang"
FROM fact_trips t
JOIN dim_operators o ON t.operator_id = o.operator_id
JOIN dim_routes r ON t.route_id = r.route_id
LIMIT 500;
"""
df_trips = pd.read_sql_query(query_trips, conn)

# 6. Sheet 6: Tahapan Delay
query_stages = """
SELECT 
    s.stage_name AS "Tahapan Perjalanan",
    COUNT(s.stage_id) AS "Jumlah Sampel Etape",
    ROUND(AVG(s.actual_duration_min), 1) AS "Rata-Rata Durasi Aktual (Menit)",
    ROUND(AVG(s.delay_variance_min), 1) AS "Kontribusi Rata-Rata Delay (Menit)",
    ROUND(AVG(s.fuel_burned_liters), 2) AS "Rata-Rata Solar Terbakar (Liter)",
    ROUND(AVG(s.idle_duration_min), 1) AS "Durasi Mesin Idle (Menit)"
FROM fact_trip_stages s
GROUP BY s.stage_name
ORDER BY "Kontribusi Rata-Rata Delay (Menit)" DESC;
"""
df_stages = pd.read_sql_query(query_stages, conn)

conn.close()

# 7. Write into Excel with multiple styled sheets
target_file = OUTPUT_EXCEL
writer = None
try:
    writer = pd.ExcelWriter(target_file, engine="openpyxl")
    writer.__enter__()
except PermissionError:
    target_file = os.path.join(BASE_DIR, "data", "nusantarabus_passenger_sentiment_complete.xlsx")
    print(f"Original file is open in another app. Writing to fallback: {target_file}")
    writer = pd.ExcelWriter(target_file, engine="openpyxl")
    writer.__enter__()

try:
    df_granular_reviews.to_excel(writer, sheet_name="1_Ulasan_Riil_Per_Penumpang", index=False)
    df_sentiment_summary.to_excel(writer, sheet_name="2_Ringkasan_Sentimen_PO", index=False)
    df_raw_scraped.to_excel(writer, sheet_name="3_Jadwal_Scraped_RedBus", index=False)
    df_operators.to_excel(writer, sheet_name="4_Peringkat_Kinerja_PO", index=False)
    df_trips.to_excel(writer, sheet_name="5_Sampel_Perjalanan_TransJawa", index=False)
    df_stages.to_excel(writer, sheet_name="6_Akar_Masalah_Delay_Etape", index=False)
finally:
    if writer:
        writer.__exit__(None, None, None)

# 8. Polish Styling
from openpyxl import load_workbook
wb = load_workbook(target_file)

header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Deep Slate 900
header_font = Font(name="Segoe UI", size=10, bold=True, color="F8FAFC")
thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

for ws in wb.worksheets:
    ws.views.sheetView[0].showGridLines = True
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ws.row_dimensions[1].height = 26

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell.border = thin_border
            cell.font = Font(name="Segoe UI", size=9.5) if cell.row > 1 else header_font
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 70)

wb.save(target_file)
print(f"Updated complete dataset with granular passenger reviews saved to: {target_file}")
